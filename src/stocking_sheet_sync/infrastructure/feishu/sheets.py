from __future__ import annotations

import logging
import re
import time
from datetime import date, datetime
from datetime import time as daytime
from io import BytesIO
from pathlib import Path
from urllib.parse import quote
from xml.etree.ElementTree import tostring

import httpx
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

from stocking_sheet_sync.domain.products import column_name
from stocking_sheet_sync.infrastructure.feishu.client import FeishuClient
from stocking_sheet_sync.settings import load_config

LOG = logging.getLogger(__name__)


def create_client() -> FeishuClient:
    """读取项目配置并创建数据应用客户端，由调用方负责关闭。"""
    config = load_config()
    return FeishuClient(config, config.feishu_data_app_id, config.feishu_data_app_secret, "sheets")


def revision(client: FeishuClient, token: str, sheet_id: str) -> int:
    """使用 client 获取 token 中 sheet_id 的当前工作簿版本号。"""
    data = client._request(
        "GET",
        f"/open-apis/sheets/v2/spreadsheets/{quote(token, safe='')}/values_batch_get",
        params={"ranges": f"{sheet_id}!A1:A1", "valueRenderOption": "Formula"},
    )
    return int(data["revision"])


def current_revision(token: str, sheet_id: str) -> int:
    """通过数据应用获取 token 中 sheet_id 的版本号，完成后关闭连接。"""
    client = create_client()
    try:
        return revision(client, token, sheet_id)
    finally:
        client.close()


def export_workbook(client: FeishuClient, token: str) -> bytes:
    """
    功能说明：通过官方导出任务获取完整 XLSX，用于读取样式及公式类型。

    参数：
        client：已配置应用身份的飞书服务端客户端。
        token：电子表格 token。

    返回值：导出文件字节；任务失败或超过两分钟抛出异常。
    """
    task = client._request(
        "POST",
        "/open-apis/drive/v1/export_tasks",
        retry=False,
        json_body={"token": token, "type": "sheet", "file_extension": "xlsx"},
    )
    LOG.debug("表格样式导出开始：token=%s", token)
    deadline = time.monotonic() + 120
    while time.monotonic() < deadline:
        result = client._request(
            "GET",
            "/open-apis/drive/v1/export_tasks/" + quote(task["ticket"], safe=""),
            params={"token": token},
        )["result"]
        if result["job_status"] == 0:
            response = client._client.get(
                "/open-apis/drive/v1/export_tasks/file/"
                + quote(result["file_token"], safe="")
                + "/download",
                headers={"Authorization": "Bearer " + client._get_access_token()},
            )
            if not response.is_success or not response.content.startswith(b"PK"):
                raise RuntimeError(f"表格导出下载失败：status={response.status_code}")
            LOG.debug("表格样式导出完成：bytes=%d", len(response.content))
            return response.content
        if result["job_status"] not in (1, 2):
            raise RuntimeError(f"表格导出任务失败：job_status={result['job_status']}")
        time.sleep(1)
    raise RuntimeError("表格样式导出超时")


def _xml(value) -> str:
    """将样式对象 value 序列化为稳定文本，独立于导出文件内的样式编号。"""
    return tostring(value.to_tree(), encoding="unicode")


def dimension_snapshot(value) -> dict:
    """提取行列维度 value 的尺寸和实际样式，使用内容比较而非导出文件内的样式编号。"""
    fields = {"r", "hidden", "outlineLevel", "collapsed", "thickTop", "thickBot", "customFormat"}
    if value.__class__.__name__ == "ColumnDimension":
        fields.update({"width", "bestFit", "min", "max", "customWidth"})
    result = {k: v for k, v in dict(value).items() if k in fields}
    result["style_definition"] = {
        "font": _xml(value.font),
        "fill": _xml(value.fill),
        "border": _xml(value.border),
        "alignment": _xml(value.alignment),
        "protection": _xml(value.protection),
        "number_format": value.number_format,
    }
    return result


def read_sheet(
    token: str,
    sheet_id: str,
    *,
    include_style: bool = True,
    archive_path: Path | None = None,
    client: FeishuClient | None = None,
) -> dict:
    """
    功能说明：通过服务端 API 读取全表数值，以 XLSX 补充公式、样式和布局并核对版本。

    参数：
        token：电子表格 token。
        sheet_id：需要读取的工作表 ID。
        include_style：是否包含样式与布局；公式类型始终从导出文件确认。
        archive_path：可选导出文件保存路径，供审阅原始证据。
        client：可选数据应用客户端；传入时由调用方关闭。

    返回值：完整单元格快照，兼容商品和销量匹配流程；任何版本冲突均终止读取。
    """
    owned = client is None
    client = client or create_client()
    workbook = None
    try:
        initial = revision(client, token, sheet_id)
        base = "/open-apis/sheets/v3/spreadsheets/" + quote(token, safe="")
        book = client._request("GET", base)["spreadsheet"]
        sheets = client._request("GET", base + "/sheets/query")["sheets"]
        sheet = next((s for s in sheets if s["sheet_id"] == sheet_id), None)
        if sheet is None or sheet.get("resource_type") != "sheet":
            raise ValueError("指定工作表不存在或不是普通电子表格")
        rows, cols = sheet["grid_properties"]["row_count"], sheet["grid_properties"]["column_count"]
        if not 1 <= rows <= 50000 or not 1 <= cols <= 100:
            raise ValueError("服务端读取支持最多50000行、100列")
        exported = export_workbook(client, token)
        if archive_path is not None:
            archive_path.parent.mkdir(parents=True, exist_ok=True)
            archive_path.write_bytes(exported)
        workbook = load_workbook(BytesIO(exported), data_only=False)
        if sheet["title"] not in workbook.sheetnames:
            raise ValueError("导出文件缺少目标工作表")
        ws = workbook[sheet["title"]]
        merges = [str(m) for m in ws.merged_cells.ranges]
        api_merges = {
            f"{column_name(m['start_column_index'] + 1)}{m['start_row_index'] + 1}:"
            f"{column_name(m['end_column_index'] + 1)}{m['end_row_index'] + 1}"
            for m in sheet.get("merges", [])
        }
        if set(merges) != api_merges:
            raise ValueError("导出合并范围与服务端元数据不一致")
        cells = {}
        chunk = max(1, min(200, 5000 // cols))
        for start in range(1, rows + 1, chunk):
            stop = min(rows, start + chunk - 1)
            area = f"{sheet_id}!A{start}:{column_name(cols)}{stop}"
            data = client._request(
                "GET",
                f"/open-apis/sheets/v2/spreadsheets/{quote(token, safe='')}/values_batch_get",
                params={"ranges": area, "valueRenderOption": "UnformattedValue"},
            )
            if data["revision"] != initial:
                raise RuntimeError("读取期间表格有修改，请重新运行")
            ranges = data["valueRanges"]
            if len(ranges) != 1 or ranges[0]["range"] != area:
                raise ValueError("服务端返回范围与请求不一致")
            values = ranges[0].get("values") or []
            if len(values) > stop - start + 1 or any(len(row) > cols for row in values):
                raise ValueError("服务端单元格矩阵超出请求范围")
            for row_no in range(start, stop + 1):
                line = values[row_no - start] if row_no - start < len(values) else []
                for col in range(1, cols + 1):
                    cell = ws.cell(row_no, col)
                    value = line[col - 1] if col <= len(line) else None
                    rich_text = None
                    if isinstance(value, list):
                        if not value or any(
                            not isinstance(segment, dict)
                            or segment.get("type") != "text"
                            or not isinstance(segment.get("text"), str)
                            for segment in value
                        ):
                            raise ValueError(f"单元格含暂不支持的复合内容：{cell.coordinate}")
                        rich_text = value
                        value = "".join(segment["text"] for segment in rich_text)
                    if cell.data_type != "f":
                        exported_value = cell.value
                        if isinstance(exported_value, (date, datetime, daytime)):
                            exported_value = to_excel(exported_value, workbook.epoch)
                        if value != exported_value and not (
                            value in (None, "") and exported_value in (None, "")
                        ):
                            raise ValueError(f"服务端值与完整导出不一致：{cell.coordinate}")
                    item = {} if value is None else {"value": value}
                    if rich_text is not None:
                        item["rich_text"] = rich_text
                    if cell.data_type == "f":
                        if not isinstance(cell.value, str):
                            raise ValueError("目标工作表含暂不支持的数组公式")
                        item["formula"] = cell.value
                    if include_style:
                        item["cell_styles"] = {
                            "font": _xml(cell.font),
                            "fill": _xml(cell.fill),
                            "alignment": _xml(cell.alignment),
                            "protection": _xml(cell.protection),
                            "number_format": cell.number_format,
                        }
                        item["border_styles"] = _xml(cell.border)
                        if cell.comment:
                            item["note"] = {
                                "text": cell.comment.text,
                                "author": cell.comment.author,
                            }
                    cells[f"{column_name(col)}{row_no}"] = item
            LOG.debug("服务端读取分页完成：sheet_id=%s rows=%s:%s", sheet_id, start, stop)
        if revision(client, token, sheet_id) != initial:
            raise RuntimeError("读取期间表格有修改，请重新运行")
        layout = {
            "revision": initial,
            "grid_properties": sheet["grid_properties"],
            "hidden": sheet["hidden"],
            "row_dimensions": {str(k): dimension_snapshot(v) for k, v in ws.row_dimensions.items()},
            "column_dimensions": {
                k: dimension_snapshot(v) for k, v in ws.column_dimensions.items()
            },
            "sheet_format": {
                key: value
                for key, value in dict(ws.sheet_format).items()
                if key
                in {
                    "baseColWidth",
                    "defaultColWidth",
                    "zeroHeight",
                    "thickTop",
                    "thickBottom",
                    "outlineLevelRow",
                    "outlineLevelCol",
                }
            },
            "freeze_panes": ws.freeze_panes,
            "data_validations": _xml(ws.data_validations),
        }
        return {
            "spreadsheet_token": token,
            "sheet_id": sheet_id,
            "title": book["title"],
            "revision": initial,
            "row_count": rows,
            "column_count": cols,
            "cells": cells,
            "merges": sorted(merges),
            "transport": "feishu_openapi",
            **({"layout": layout, "sheet_metadata": sheet} if include_style else {}),
        }
    finally:
        if workbook is not None:
            workbook.close()
        if owned:
            client.close()


def write_sales_ranges(
    token: str,
    sheet_id: str,
    operations: list[dict],
    *,
    expected_revision: int,
    overwrite_cells: set[str] | None = None,
    client: FeishuClient | None = None,
) -> dict:
    """
    功能说明：复核版本与目标空白状态后，以一次服务端请求写入多个销量范围。

    参数：
        token：目标电子表格 token。
        sheet_id：目标工作表 ID。
        operations：原生 valueRanges 数组，包含整数或单平台 SUM 合计公式。
        expected_revision：读取并核对过的工作簿版本。
        overwrite_cells：允许覆盖已有非公式值的单元格集合；默认不覆盖。
        client：可选数据应用客户端；传入时由调用方关闭。

    返回值：服务端批量写入结果；写请求仅发送一次，异常时必须回读确认。
    """
    if not operations:
        raise ValueError("不能提交空写入请求")
    claimed = set()
    for item in operations:
        if not item["range"].startswith(sheet_id + "!"):
            raise ValueError("写入范围不属于目标工作表")
        area = item["range"].split("!", 1)[1]
        bounds = re.fullmatch(r"([A-Z]+)([1-9][0-9]*):\1([1-9][0-9]*)", area)
        if not bounds:
            raise ValueError("销量写入范围必须是单列矩形")
        col, first, last = bounds.groups()
        addresses = {f"{col}{r}" for r in range(int(first), int(last) + 1)}
        if (
            not addresses
            or len(item["values"]) != len(addresses)
            or any(len(row) != 1 for row in item["values"])
            or claimed & addresses
        ):
            raise ValueError("销量写入范围重叠或矩阵尺寸不匹配")
        claimed.update(addresses)
        for offset, row in enumerate(item["values"]):
            value = row[0]
            if type(value) is int and value >= 0:
                continue
            match = (
                re.fullmatch(
                    r"=SUM\(([A-Z]+)([1-9][0-9]*):\1([1-9][0-9]*)\)", value.get("text", "")
                )
                if isinstance(value, dict)
                else None
            )
            if (
                not match
                or value.get("type") != "formula"
                or match.group(1) != col
                or not 1 <= int(match.group(2)) <= int(match.group(3)) < int(first) + offset
            ):
                raise ValueError("仅支持非负整件数或引用本列上方范围的SUM合计公式")
    owned = client is None
    client = client or create_client()
    try:
        path = f"/open-apis/sheets/v2/spreadsheets/{quote(token, safe='')}"
        current = client._request(
            "GET",
            path + "/values_batch_get",
            params={
                "ranges": ",".join(item["range"] for item in operations),
                "valueRenderOption": "Formula",
            },
        )
        if current["revision"] != expected_revision:
            raise ValueError("提交前表格版本发生变化，请重新预览")
        if [r["range"] for r in current["valueRanges"]] != [o["range"] for o in operations]:
            raise ValueError("写入前目标范围回读不完整")
        for block in current["valueRanges"]:
            area = block["range"].split("!", 1)[1]
            col, first = re.match(r"([A-Z]+)([0-9]+)", area).groups()
            for offset, row in enumerate(block.get("values", [])):
                for value in row:
                    address = f"{col}{int(first) + offset}"
                    if value not in (None, "") and (
                        address not in (overwrite_cells or set())
                        or isinstance(value, (dict, list))
                        or isinstance(value, str)
                        and value.startswith("=")
                    ):
                        raise ValueError("写入前目标单元格已有内容，请重新核对")
        LOG.debug("开始服务端销量写入：ranges=%d revision=%d", len(operations), expected_revision)
        return client._request(
            "POST",
            path + "/values_batch_update",
            retry=False,
            json_body={"valueRanges": operations},
        )
    except httpx.TransportError:
        raise RuntimeError("飞书写入连接异常，结果不确定；请回读表格，不要直接重试") from None
    finally:
        if owned:
            client.close()


def apply_update(
    token: str,
    sheet_id: str,
    update: dict,
    expected_revision: int,
    *,
    client=None,
    on_progress=None,
) -> dict:
    """
    功能说明：通过飞书服务端接口顺序执行补列，逐步保存结果并检查版本。

    参数：
        token：目标副本 token。
        sheet_id：工作表 ID。
        update：由 build_update 生成的请求计划。
        expected_revision：执行前已审阅的版本。
        client：可选数据应用客户端；传入时由调用方关闭。
        on_progress：可选回调，每次请求前后保存执行日志。
    返回值：各步骤结果；任一步失败立即停止，已提交步骤不自动重试或回滚。
    """
    owned = client is None
    client = client or create_client()
    journal = {"status": "running", "steps": [], "revision": expected_revision}
    base = f"/open-apis/sheets/v2/spreadsheets/{quote(token, safe='')}"
    try:
        for index, operation in enumerate(update["operations"]):
            if revision(client, token, sheet_id) != journal["revision"]:
                raise ValueError("补列期间版本发生变化，请回读核对，不要直接重试")
            step = {"index": index, "operation": operation, "status": "sending"}
            journal["steps"].append(step)
            if on_progress:
                on_progress(journal)
            LOG.debug("补列开始：token=%s step=%d endpoint=%s", token, index, operation["endpoint"])
            step["response"] = client._request(
                operation["method"],
                base + "/" + operation["endpoint"],
                json_body=operation["body"],
                retry=False,
            )
            step["status"] = "acknowledged"
            # 部分结构接口不返回版本；下一步前再次读取，最终以全表核验确认。
            journal["revision"] = revision(client, token, sheet_id)
            if on_progress:
                on_progress(journal)
        journal["status"] = "submitted"
        return journal
    finally:
        if owned:
            client.close()
