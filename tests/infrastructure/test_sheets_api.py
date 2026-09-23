from io import BytesIO
from unittest.mock import Mock

import httpx
import pytest
from openpyxl import Workbook
from openpyxl.styles import Border, Side

from stocking_sheet_sync.infrastructure.feishu.sheets import read_sheet, write_sales_ranges


def transport_client(monkeypatch, *, nonempty=False, current=1, timeout=False):
    client = Mock()
    requests = []

    def request(method, path, **kwargs):
        requests.append((method, path, kwargs))
        if method == "GET":
            return {
                "revision": current,
                "valueRanges": [
                    {"range": "test!N4:N5", "values": [[0 if nonempty else None], [None]]}
                ],
            }
        if timeout:
            raise httpx.ReadTimeout("网络中断")
        return {"revision": 2, "responses": [{"updatedCells": 2}]}

    client._request.side_effect = request
    monkeypatch.setattr(
        "stocking_sheet_sync.infrastructure.feishu.sheets.create_client", lambda: client
    )
    return client, requests


def test_native_write_sends_numeric_values_once_and_preserves_style(monkeypatch):
    client, requests = transport_client(monkeypatch)
    ops = [{"range": "test!N4:N5", "values": [[0], [7]]}]
    result = write_sales_ranges("token", "test", ops, expected_revision=1)
    assert result["revision"] == 2
    assert requests[-1] == (
        "POST",
        "/open-apis/sheets/v2/spreadsheets/token/values_batch_update",
        {"retry": False, "json_body": {"valueRanges": ops}},
    )
    assert len(requests) == 2
    client.close.assert_called_once()


@pytest.mark.parametrize("nonempty,current", [(True, 1), (False, 2)])
def test_native_preflight_refuses_existing_zero_and_revision_conflict(
    monkeypatch, nonempty, current
):
    client, requests = transport_client(monkeypatch, nonempty=nonempty, current=current)
    with pytest.raises(ValueError):
        write_sales_ranges(
            "token", "test", [{"range": "test!N4:N5", "values": [[0], [7]]}], expected_revision=1
        )
    assert len(requests) == 1
    client.close.assert_called_once()


def test_native_write_uncertain_network_result_is_not_retried(monkeypatch):
    _, requests = transport_client(monkeypatch, timeout=True)
    with pytest.raises(RuntimeError, match="结果不确定"):
        write_sales_ranges(
            "token", "test", [{"range": "test!N4:N5", "values": [[0], [7]]}], expected_revision=1
        )
    assert sum(m == "POST" for m, _, _ in requests) == 1


@pytest.mark.parametrize("rich_text", [False, True])
@pytest.mark.parametrize("missing_value", [False, True])
def test_native_read_retains_formulas_styles_and_all_empty_coordinates(
    monkeypatch, missing_value, rich_text
):
    book = Workbook()
    sheet = book.active
    sheet.title = "样本"
    sheet["A1"] = "00123"
    sheet["B1"] = "=1+1"
    sheet["A1"].border = Border(left=Side(style="thin", color="001F2329"))
    sheet.row_dimensions[1].hidden = True
    if missing_value:
        sheet["A3"] = "服务端遗漏的内容"
    stream = BytesIO()
    book.save(stream)
    book.close()
    client = Mock()
    segments = [
        {"type": "text", "text": "001", "segmentStyle": {"bold": False}},
        {"type": "text", "text": "23", "segmentStyle": {"bold": True}},
    ]

    def request(method, path, **kwargs):
        if path.endswith("/sheets/query"):
            return {
                "sheets": [
                    {
                        "sheet_id": "test",
                        "title": "样本",
                        "resource_type": "sheet",
                        "hidden": False,
                        "grid_properties": {"row_count": 3, "column_count": 2},
                    }
                ]
            }
        if path.endswith("/values_batch_get"):
            area = kwargs["params"]["ranges"]
            return {
                "revision": 1,
                "valueRanges": [
                    {"range": area, "values": [[segments if rich_text else "00123", 2]]}
                ],
            }
        return {"spreadsheet": {"title": "样本"}}

    client._request.side_effect = request
    monkeypatch.setattr(
        "stocking_sheet_sync.infrastructure.feishu.sheets.create_client", lambda: client
    )
    monkeypatch.setattr(
        "stocking_sheet_sync.infrastructure.feishu.sheets.export_workbook",
        lambda *a: stream.getvalue(),
    )
    if missing_value:
        with pytest.raises(ValueError, match="完整导出不一致"):
            read_sheet("token", "test")
        return
    data = read_sheet("token", "test")
    assert len(data["cells"]) == 6
    assert data["cells"]["A1"]["value"] == "00123"
    if rich_text:
        assert data["cells"]["A1"]["rich_text"] == segments
    assert data["cells"]["B1"]["formula"] == "=1+1"
    assert data["cells"]["B1"]["value"] == 2
    assert "value" not in data["cells"]["B3"]
    assert "001F2329" in data["cells"]["A1"]["border_styles"]
    assert data["layout"]["row_dimensions"]["1"]["hidden"] == "1"
    client.close.assert_called_once()


@pytest.mark.parametrize(
    "operations",
    [
        [{"range": "test!N4:N5", "values": [[1], [2], [3]]}],
        [{"range": "test!N4:O4", "values": [[1, 2]]}],
        [{"range": "test!N4:N4", "values": [[1]]}, {"range": "test!N4:N4", "values": [[2]]}],
    ],
)
def test_native_writer_rejects_oversized_or_overlapping_ranges_before_access(
    monkeypatch, operations
):
    create = Mock(side_effect=AssertionError("无效范围不应访问飞书"))
    monkeypatch.setattr("stocking_sheet_sync.infrastructure.feishu.sheets.create_client", create)
    with pytest.raises(ValueError):
        write_sales_ranges("token", "test", operations, expected_revision=1)
    create.assert_not_called()


def test_native_writer_sends_typed_formula_and_rejects_circular_total(monkeypatch):
    _, requests = transport_client(monkeypatch)
    formula = {"type": "formula", "text": "=SUM(N1:N3)"}
    operations = [{"range": "test!N4:N5", "values": [[formula], [7]]}]
    write_sales_ranges("token", "test", operations, expected_revision=1)
    assert requests[-1][2]["json_body"]["valueRanges"][0]["values"][0][0] == formula
    formula["text"] = "=SUM(N1:N4)"
    with pytest.raises(ValueError, match="SUM"):
        write_sales_ranges("token", "test", operations, expected_revision=1)


def test_dimension_snapshot_compares_style_definitions_instead_of_export_ids():
    from openpyxl.styles import Font

    from stocking_sheet_sync.infrastructure.feishu.sheets import dimension_snapshot

    first, second = Workbook(), Workbook()
    first.active.row_dimensions[1].font = Font(name="宋体", size=14)
    second.active["A1"].font = Font(name="Arial", size=10)
    second.active.row_dimensions[1].font = Font(name="宋体", size=14)
    left, right = first.active.row_dimensions[1], second.active.row_dimensions[1]
    assert left._style.fontId != right._style.fontId
    assert dimension_snapshot(left) == dimension_snapshot(right)
    right.font = Font(name="宋体", size=16)
    assert dimension_snapshot(left) != dimension_snapshot(right)


@pytest.mark.parametrize("allowed,revision", [(True, 1), (False, 1), (True, 2)])
def test_overwrite_requires_explicit_cell_and_matching_revision(monkeypatch, allowed, revision):
    _, calls = transport_client(monkeypatch, nonempty=True, current=revision)
    args = dict(expected_revision=1, overwrite_cells={"N4"} if allowed else {"N5"})
    operations = [{"range": "test!N4:N5", "values": [[20], [7]]}]
    if allowed and revision == 1:
        write_sales_ranges("token", "test", operations, **args)
        assert calls[-1][0] == "POST"
    else:
        with pytest.raises(ValueError):
            write_sales_ranges("token", "test", operations, **args)
        assert len(calls) == 1
